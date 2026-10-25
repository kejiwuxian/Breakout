#!/usr/bin/env python3
"""Render SpaceX-themed pitch slides (1920×1080 PNG) + PPTX from script + shots."""

from __future__ import annotations

import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parent
ART = Path("/opt/cursor/artifacts/presentation")
SHOTS = ART / "shots"
SLIDES = ART / "slides"
WORK = ART / "work"
SCRIPT = json.loads((ROOT / "script.json").read_text())

W, H = 1920, 1080
BLACK = (0, 0, 0)
NEAR = (10, 10, 10)
PANEL = (20, 20, 20)
WHITE = (255, 255, 255)
DIM = (160, 160, 160)
MUTED = (110, 110, 110)
RED = (227, 28, 35)
HAIR = (40, 40, 40)


def font(size: int, mono: bool = False, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = []
    if mono:
        candidates = [
            "/usr/share/fonts/truetype/jetbrains-mono/JetBrainsMono-Regular.ttf",
            "/usr/share/fonts/truetype/jetbrains-mono/JetBrainsMonoNL-Regular.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf",
        ]
    elif bold:
        candidates = [
            "/usr/share/fonts/truetype/macos/Inter-Bold.ttf",
            "/usr/share/fonts/truetype/macos/Inter-SemiBold.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        ]
    else:
        candidates = [
            "/usr/share/fonts/truetype/macos/Inter-Regular.ttf",
            "/usr/share/fonts/truetype/macos/Inter-Medium.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        ]
    for p in candidates:
        if Path(p).is_file():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def find_font_files() -> None:
    import subprocess

    out = subprocess.check_output(["fc-list", ":file"], text=True)
    inter = [l.split(":")[0] for l in out.splitlines() if "Inter" in l and l.endswith((".ttf", ".otf"))]
    print("inter fonts sample", inter[:8])


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_w: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur = ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textlength(trial, font=fnt) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def base_canvas() -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (W, H), BLACK)
    draw = ImageDraw.Draw(img)
    # top hairline + LIVE pill chrome
    draw.rectangle([0, 0, W, 4], fill=RED)
    draw.rectangle([0, H - 48, W, H], fill=NEAR)
    draw.rectangle([48, H - 36, 48 + 70, H - 14], fill=RED)
    f_live = font(14, mono=True, bold=True)
    draw.text((58, H - 34), "LIVE", fill=WHITE, font=f_live)
    f_foot = font(16, mono=True)
    draw.text((140, H - 34), "NEUTRINO FOLLOW-UP  ·  PROJECT SUPERNOVA  ·  SPACEX TRACK", fill=DIM, font=f_foot)
    return img, draw


def draw_eyebrow(draw: ImageDraw.ImageDraw, text: str, x: int = 80, y: int = 56) -> None:
    f = font(18, mono=True)
    draw.text((x, y), text.upper(), fill=RED, font=f)


def fit_shot(path: Path, box: tuple[int, int, int, int]) -> Image.Image:
    x0, y0, x1, y1 = box
    tw, th = x1 - x0, y1 - y0
    src = Image.open(path).convert("RGB")
    # cover-fit
    scale = max(tw / src.width, th / src.height)
    nw, nh = int(src.width * scale), int(src.height * scale)
    src = src.resize((nw, nh), Image.Resampling.LANCZOS)
    left = (nw - tw) // 2
    top = (nh - th) // 2
    return src.crop((left, top, left + tw, top + th))


def render_title(slide: dict) -> Image.Image:
    img, draw = base_canvas()
    draw_eyebrow(draw, slide["eyebrow"])
    f_h = font(96, bold=True)
    f_s = font(36)
    f_f = font(22, mono=True)
    # centered brand block
    lines = wrap(draw, slide["headline"], f_h, W - 160)
    y = 320
    for line in lines:
        w = draw.textlength(line, font=f_h)
        draw.text(((W - w) / 2, y), line, fill=WHITE, font=f_h)
        y += 110
    sub = slide.get("sub", "")
    w = draw.textlength(sub, font=f_s)
    draw.text(((W - w) / 2, y + 20), sub, fill=DIM, font=f_s)
    foot = slide.get("footer", "")
    w = draw.textlength(foot, font=f_f)
    draw.text(((W - w) / 2, H - 120), foot, fill=MUTED, font=f_f)
    # side accent bars
    draw.rectangle([80, 300, 84, 520], fill=RED)
    draw.rectangle([W - 84, 300, W - 80, 520], fill=RED)
    return img


def render_bullets(slide: dict) -> Image.Image:
    img, draw = base_canvas()
    draw_eyebrow(draw, slide["eyebrow"])
    f_h = font(56, bold=True)
    f_b = font(28)
    y = 120
    for line in wrap(draw, slide["headline"], f_h, W - 200):
        draw.text((80, y), line, fill=WHITE, font=f_h)
        y += 68
    y += 36
    draw.rectangle([80, y, 200, y + 3], fill=RED)
    y += 40
    for b in slide["bullets"]:
        # red dash marker
        draw.rectangle([80, y + 18, 110, y + 22], fill=RED)
        for line in wrap(draw, b, f_b, W - 240):
            draw.text((130, y), line, fill=DIM if y > 0 else WHITE, font=f_b)
            # use near-white for body
            draw.text((130, y), line, fill=(210, 210, 210), font=f_b)
            y += 42
        y += 22
    return img


def render_shot(slide: dict) -> Image.Image:
    img, draw = base_canvas()
    draw_eyebrow(draw, slide["eyebrow"])
    f_h = font(44, bold=True)
    f_c = font(22)
    y = 100
    for line in wrap(draw, slide["headline"], f_h, 900):
        draw.text((80, y), line, fill=WHITE, font=f_h)
        y += 54
    # left caption column
    cy = y + 24
    draw.rectangle([80, cy, 160, cy + 3], fill=RED)
    cy += 24
    for line in wrap(draw, slide.get("caption", ""), f_c, 520):
        draw.text((80, cy), line, fill=DIM, font=f_c)
        cy += 32

    # screenshot frame on right
    shot_path = SHOTS / slide["shot"]
    if not shot_path.is_file():
        # try alternate early names
        alt = SHOTS / slide["shot"].replace("spacex-demo-", "spacex-")
        shot_path = alt if alt.is_file() else shot_path
    frame = (640, 140, W - 64, H - 80)
    draw.rectangle([frame[0] - 2, frame[1] - 2, frame[2] + 2, frame[3] + 2], outline=HAIR, width=2)
    if shot_path.is_file():
        cropped = fit_shot(shot_path, frame)
        img.paste(cropped, (frame[0], frame[1]))
    else:
        draw.rectangle(frame, fill=PANEL)
        miss = font(24, mono=True)
        draw.text((frame[0] + 40, frame[1] + 40), f"MISSING\n{slide['shot']}", fill=RED, font=miss)

    # small LIVE tag over frame
    draw.rectangle([frame[0] + 16, frame[1] + 16, frame[0] + 86, frame[1] + 40], fill=RED)
    draw.text((frame[0] + 26, frame[1] + 18), "LIVE", fill=WHITE, font=font(14, mono=True, bold=True))
    return img


def render_all() -> list[Path]:
    SLIDES.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for i, slide in enumerate(SCRIPT["slides"], start=1):
        kind = slide["kind"]
        if kind == "title":
            img = render_title(slide)
        elif kind == "bullets":
            img = render_bullets(slide)
        else:
            img = render_shot(slide)
        out = SLIDES / f"{i:02d}-{slide['id']}.png"
        img.save(out, "PNG", optimize=True)
        paths.append(out)
        print("slide", out)
    return paths


def build_pptx(paths: list[Path]) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for p in paths:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(p), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
    dest = ART / "Project_Supernova_SpaceX_Pitch.pptx"
    # also copy into repo presentation folder
    repo_dest = ROOT / "Project_Supernova_SpaceX_Pitch.pptx"
    prs.save(str(dest))
    prs.save(str(repo_dest))
    print("pptx", dest)
    return dest


def main() -> None:
    find_font_files()
    paths = render_all()
    build_pptx(paths)
    # copy slides into repo for review (pngs may be large; keep under presentation/spacex/slides)
    repo_slides = ROOT / "slides"
    repo_slides.mkdir(exist_ok=True)
    for p in paths:
        (repo_slides / p.name).write_bytes(p.read_bytes())
    print("done", len(paths), "slides")


if __name__ == "__main__":
    main()
